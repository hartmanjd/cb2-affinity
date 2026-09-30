"""Build one readable SVR workbook for baseline and tuning, without fitting."""
from pathlib import Path
import csv
import json
import math
import os
from statistics import mean, pstdev
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter
from build_mlp_summary import digest, read_json, read_csv, flatten, numeric
from forest_diagnostics import derive

ROOT=Path(__file__).resolve().parents[1]
SVR=Path('provenance/models/support_vector_regression')
TABLES={'metrics.csv':'Scores by Split','training_summary.csv':'Training Summary',
        'validation_subgroups.csv':'Saved Similarity','predictions.csv':'Predictions',
        'baseline_comparison.csv':'Baseline Comparison','cv_folds.csv':'CV Folds','cv_results.csv':'Search Ranking'}


def build_summary(root=ROOT):
    root=Path(root);parent=root/SVR;files=parent/'files'
    manifests={v:read_json(files/v/'manifest.json') for v in ['baseline','tuning']}
    baseline,tuning=manifests['baseline'],manifests['tuning']
    assert baseline.get('completed_at_utc') and tuning.get('completed_at_utc')
    assert tuning['baseline_manifest']==(SVR/'files/baseline/manifest.json').as_posix()
    assert digest(root/tuning['baseline_manifest'])==tuning['baseline_manifest_sha256']
    assert baseline['source_preparation_manifest']==tuning['source_preparation_manifest']
    prep_path=root/baseline['source_preparation_manifest'];prep=read_json(prep_path)
    diag=read_json(files/'diagnostics/manifest.json')
    for m in [baseline,tuning,diag]:
        assert m['source_preparation_manifest']==baseline['source_preparation_manifest']
        assert m['source_preparation_manifest_sha256']==digest(prep_path)
        for field in ['output_sha256','source_artifact_sha256']:
            for name,expected in m.get(field,{}).items():
                assert digest(root/name)==expected, f'Changed evidence: {name}'
    metrics={v:read_csv(files/v/'metrics.csv') for v in manifests}
    neighbors_path=root/prep['output_folder']/'heldout_training_neighbors.csv'
    neighbors=read_csv(neighbors_path)
    assert digest(neighbors_path)==prep['artifact_sha256'][neighbors_path.relative_to(root).as_posix()]
    pairs=read_csv(files/'diagnostics/cliff_pairs.csv')
    predictions={v:read_csv(files/v/'predictions.csv') for v in manifests}
    derived={v:derive(predictions[v],pairs,neighbors) for v in manifests}
    # Confirm saved search means and the selected candidates independently.
    candidates=read_csv(files/'tuning/cv_results.csv')
    folds=int(tuning['configuration']['folds'])
    for split in ['random','scaffold']:
        rows=[r for r in candidates if r['split_strategy']==split]
        for r in rows:
            scores=[float(r[f'fold_{i}_validation_mae_pki']) for i in range(1,folds+1)]
            assert math.isclose(mean(scores),float(r['mean_cv_mae_pki']),abs_tol=1e-12)
            assert math.isclose(pstdev(scores),float(r['std_cv_mae_pki']),abs_tol=1e-12)
        best=min(rows,key=lambda r:(float(r['mean_cv_mae_pki']),int(r['candidate_index'])))
        assert int(best['candidate_index'])==tuning['selection'][split]['best_candidate_index']
    book=Workbook(write_only=True);counts={}
    def cell(ws,value,style=None,link=None):
        c=WriteOnlyCell(ws,value=value)
        if isinstance(value,str):c.data_type='s'
        c.font=Font(name='Calibri',size=19 if style=='title' else 11,bold=bool(style),color='FFFFFF' if style else '18354A')
        c.alignment=Alignment(vertical='top',wrap_text=True)
        if style:c.fill=PatternFill('solid',fgColor='18354A' if style=='title' else '007F82')
        if isinstance(value,float):c.number_format='0.0000'
        if link:
            c.hyperlink=Hyperlink(ref=c.coordinate,location=link[1:]) if link.startswith('#') else link
            c.font=Font(name='Calibri',size=11,color='007F82',underline='single')
        return c
    def worksheet(name,widths):
        ws=book.create_sheet(name);ws.sheet_view.showGridLines=False;ws.sheet_properties.tabColor='007F82'
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        return ws
    def table(name,subtitle,headers,rows,widths=None):
        ws=worksheet(name,widths or [max(19,min(58,len(h)+3)) for h in headers]);ws.freeze_panes='A5'
        for row in [1,2]:ws.merged_cells.add(f'A{row}:{get_column_letter(len(headers))}{row}')
        ws.row_dimensions[1].height=32;ws.row_dimensions[2].height=42;ws.row_dimensions[4].height=40
        ws.append([cell(ws,name,'title')]);ws.append([cell(ws,subtitle)])
        ws.append([cell(ws,'Back to overview',link="#'Overview'!A1")]);ws.append([cell(ws,h,'header') for h in headers])
        for i,row in enumerate(rows,5):
            if name in ['Training Guide','Settings']:ws.row_dimensions[i].height=44
            ws.append([cell(ws,v) for v in row])
        ws.auto_filter.ref=f'A4:{get_column_letter(len(headers))}{4+len(rows)}';counts[name]=len(rows)
        return ws
    cover=worksheet('Overview',[19,20,18,17,19,19,19,27]);cover.freeze_panes='D6'
    def line(values,style=None,height=30,links=None):
        line.row+=1;cover.row_dimensions[line.row].height=height
        if len(values)<8:
            first='A' if len(values)==1 else 'B' if len(values)==2 else 'C'
            cover.merged_cells.add(f'{first}{line.row}:H{line.row}')
        cover.append([cell(cover,v,style,(links or {}).get(i)) for i,v in enumerate(values)])
    line.row=0
    line(['SVR | BASELINE + TUNED'],'title',38)
    line(['Saved executions',f"Baseline: {baseline['completed_at_utc']}  |  Tuned: {tuning['completed_at_utc']}"],height=36)
    line(['One deterministic SVR fit per split and variant. SVR has no random_state parameter; seed variability was not measured.'])
    line(['Validation estimates generalization; training describes fit. Reserved test sets remain unevaluated.'])
    line(['Split','Subset','Variant','Molecules','MAE (pKi)','RMSE (pKi)','R²','R² status'],'header')
    for subset in ['validation','train']:
        for split in ['random','scaffold']:
            for variant in ['baseline','tuning']:
                r=next(r for r in metrics[variant] if r['split_strategy']==split and r['subset']==subset)
                line([split.title(),subset.title(),'Tuned' if variant=='tuning' else 'Baseline',int(r['n_structures']),
                      float(r['mae_pki']),float(r['rmse_pki']),float(r['r2']) if r['r2'] else None,r['r2_status']])
    line(['Lower MAE / RMSE is better; higher R² is better. A tuned model is not necessarily better on validation.'])
    line(['MODEL AND TRAINING AT A GLANCE'],'header')
    fp=prep['configuration']['fingerprint'];p=baseline['svr_parameters']['random']
    line(['Architecture','Support vector regression (RBF)',"An RBF kernel compares fingerprint vectors. Support vectors define one continuous pKi prediction; epsilon sets the error-insensitive tube."],height=48)
    line(['Inputs','Morgan fingerprints',f"{fp['bits']} bits; radius {fp['radius']}; chirality {fp['include_chirality']}"],height=38)
    line(['Baseline','Fixed before fitting',f"kernel = {p['kernel']}; C = {p['C']}; epsilon = {p['epsilon']}; gamma = {p['gamma']}"],height=42)
    for split in ['random','scaffold']:
        selected=tuning['selection'][split]
        line([split.title()+' tuned','Selected parameters',json.dumps(selected['best_parameters'])],height=46)
        line([split.title()+' CV','Mean fold MAE',f"{selected['best_mean_cv_mae_pki']:.4f} pKi (training-only selection score)"],height=36)
    line(['Tuning',f"{tuning['counts']['candidates_per_protocol']} candidates / split",f"{folds}-fold CV: KFold for random; GroupKFold for scaffold. Refit selected settings on all training molecules."],height=46)
    line(['Dataset',f"{prep['counts']['structures']:,} structures",f"{prep['counts']['retained_measurements']:,} measurements; frozen random and scaffold splits."],height=36)
    line(['Scope','Development validation','Prior development inspected validation labels; final claims still require reserved test evaluation.'],height=44)
    line(['READ NEXT'],'header')
    for name,text in [('Training Guide','Architecture, parameter meanings, data sizes, and training procedure.'),
                      ('Search Ranking','Every candidate and its training-only CV scores.'),
                      ('Baseline Comparison','Exact saved baseline-versus-tuned results.'),
                      ('Similarity Summary','Errors by nearest-training similarity; derived from saved predictions.'),
                      ('Cliff Summary','Derived errors on frozen validation activity-cliff pairs.'),
                      ('Evidence Files','Original files, checksums, and source-table locations.')]:
        line([name,text],links={0:f"#'{name}'!A1"},height=34)
    cover.sheet_properties.pageSetUpPr.fitToPage=True;cover.page_setup.orientation='landscape';cover.page_setup.paperSize='8'
    cover.page_setup.fitToWidth=1;cover.page_setup.fitToHeight=1;cover.print_area=f'A1:H{line.row}'
    guide=[['Both','Algorithm','SVR with RBF kernel','Fits a kernel regression function to binary fingerprints. Deviations within epsilon incur no loss; C controls the penalty outside that tube.'],
           ['Both','Target','pKi',prep['target_policy']+'; higher pKi means stronger affinity.'],
           ['Baseline','Fit','Fixed settings','One deterministic SVR per outer split, fitted on its training molecules.'],
           ['Tuned','Search','GridSearchCV',tuning['configuration']['selection_rule']],
           ['Tuned','Random CV','KFold',json.dumps(tuning['configuration']['random_cv'])],
           ['Tuned','Scaffold CV','GroupKFold',json.dumps(tuning['configuration']['scaffold_cv'])],
           ['Tuned','Search grid','Declared candidates',json.dumps(tuning['configuration']['parameter_grid'])],
           ['Tuned','Final fit','Refit selected settings','All outer-training molecules; validation excluded from model selection.'],
           ['Both','MAE','Mean absolute error','Typical absolute prediction error in pKi; lower is better.'],
           ['Both','RMSE','Root mean squared error','Emphasizes large errors; lower is better.'],
           ['Both','R²','Coefficient of determination','1 is perfect; zero matches the evaluated-subset mean; negative values are possible.'],
           ['Both','Cliffs','Derived diagnostics','Unordered validation pairs with Tanimoto ≥ 0.8 and absolute observed pKi gap ≥ 1. Predicted ties count as incorrect directions.'],
           ['Both','Cliff sources','Independent frozen pair set','files/diagnostics retains the pair set for this preparation. No ongoing MLP dependency.'],
           ['Both','Similarity','Derived diagnostics','Saved nearest-training similarities joined to each variant’s validation predictions.'],
           ['Both','Models','Compressed binaries','models.joblib.gz files are preserved and linked; this report does not load or retrain them.']]
    guide.extend([
        ['Both','Input scaling','Unscaled binary Morgan bits','All 2,048 inputs share the 0/1 scale. No fitted feature scaler is used.'],
        ['Both','Reproducibility','Deterministic SVR solver','SVR has no random_state parameter. Random-protocol KFold shuffling uses random_state 42; scaffold GroupKFold is unshuffled.'],
        ['Both','Support vectors','Training structures defining the fitted function','Baseline support-vector counts are in Training Summary. They are not hidden layers or neurons.'],
        ['Both','Kernel similarity','RBF on fingerprint vectors','The fitted kernel differs from the Tanimoto similarity used for report-only neighbor and cliff diagnostics.'],
    ])
    explanations={'C':'Penalty strength for deviations outside the epsilon tube; larger values penalize them more.',
                  'epsilon':'Width of the error-insensitive tube in pKi units.',
                  'gamma':'RBF kernel coefficient: scale uses 1 / (feature count × training-input variance); auto uses 1 / feature count.',
                  'kernel':'Kernel function; both variants use the radial basis function (rbf).',
                  'cache_size':'Kernel-cache memory limit in MB.', 'tol':'Optimization stopping tolerance.',
                  'max_iter':'Solver iteration limit; -1 means no fixed limit.',
                  'shrinking':'Whether the solver uses its shrinking heuristic.',
                  'degree':'Polynomial-kernel degree; ignored by the RBF kernel.',
                  'coef0':'Constant term for polynomial/sigmoid kernels; ignored by RBF.',
                  'verbose':'Whether the underlying solver prints progress.'}
    for variant in manifests:
        for split in ['random','scaffold']:
            params=baseline['svr_parameters'][split] if variant=='baseline' else tuning['selection'][split]['effective_parameters']
            for key,value in params.items():guide.append([variant,split+' / '+key,json.dumps(value),explanations.get(key,'Saved scikit-learn estimator setting.')])
    for split,subsets in prep['counts']['split_subsets'].items():
        for subset,sizes in subsets.items():guide.append(['Both','Dataset sizes',split+' / '+subset,sizes['structures']])
    table('Training Guide','Plain-English descriptions and complete fitted-model parameters.',['Variant','Topic','Choice / value','Explanation'],guide,[18,38,62,100])
    settings=[]
    for path in sorted(files.rglob('*.json')):
        settings.extend([path.relative_to(files).as_posix(),key,value] for key,value in flatten(read_json(path)))
    settings.extend(['preparation/manifest.json',key,value] for key,value in flatten(prep))
    table('Settings','Saved JSON values, software versions and verification records.',['Source','Setting','Value'],settings,[38,75,110])
    # Combine equivalent tables while retaining variant labels and full precision.
    grouped={}
    for variant in manifests:
        for path in sorted((files/variant).glob('*.csv')):
            assert path.name in TABLES,path
            name=TABLES[path.name];rows=read_csv(path)
            bucket=grouped.setdefault(name,{'headers':['variant'],'rows':[]})
            for key in rows[0] if rows else []:
                if key not in bucket['headers']:bucket['headers'].append(key)
            bucket['rows'].extend({'variant':variant,**r} for r in rows)
    for name in ['Scores by Split','Training Summary','Baseline Comparison','Search Ranking','CV Folds','Saved Similarity','Predictions']:
        data=grouped[name];headers=data['headers']
        table(name,'Exact saved CSV data. Variant identifies the producing execution; empty fields were not saved for that variant.',headers,
              [[numeric(str(r.get(h,''))) for h in headers] for r in data['rows']])
    selected=[]
    for split,s in tuning['selection'].items():
        selected.append([split,s['best_candidate_index'],json.dumps(s['best_parameters']),s['best_mean_cv_mae_pki'],s['training_structures'],s['search_seconds'],s['refit_seconds']])
    table('Selected Parameters','Selections from the saved tuning manifest. CV scores use training folds only.',
          ['Split','Candidate','Parameters','CV MAE (pKi)','Training molecules','Search seconds','Refit seconds'],selected,[18,16,85,22,22,20,20])
    fallback={'Cliff Pair Predictions':['split_strategy','left_rdkit_smiles','right_rdkit_smiles','tanimoto','observed_difference_pki','predicted_difference_pki','absolute_difference_error_pki','correct_direction']}
    for name in ['Similarity Summary','Cliff Summary','Cliff Molecule Summary','Cliff Pair Predictions']:
        rows=[{'variant':v,**r} for v in manifests for r in derived[v][name]]
        headers=list(rows[0]) if rows else ['variant']+fallback[name]
        table(name,'DERIVED from saved predictions and frozen preparation/pair records. These are report calculations, not extra model fits.',headers,
              [[r.get(h) for h in headers] for r in rows])
    # Evidence is linked relative to the workbook, keeping links portable with this folder.
    evidence=list(sorted(p for p in files.rglob('*') if p.is_file()))+[prep_path,neighbors_path]
    ws=table('Evidence Files','Every backing file, plus preparation and neighbor sources. Binaries remain in files/.',
             ['File','Purpose','Workbook sheet','Bytes','SHA-256'],[],[62,62,30,16,72])
    for path in evidence:
        tab=TABLES.get(path.name,'Settings' if path.suffix=='.json' else 'Cliff Pair Predictions' if path.name=='cliff_pairs.csv' else 'Linked file only')
        target=Path(os.path.relpath(path,parent)).as_posix()
        purpose='Fitted models; compressed binary' if path.name.endswith('.gz') else 'Frozen cliff pair source' if path.name=='cliff_pairs.csv' else 'Original saved evidence'
        row=[target,purpose,tab,path.stat().st_size,digest(path)]
        ws.append([cell(ws,v,link=target if i==0 else f"#'{v}'!A1" if i==2 and v!='Linked file only' else None) for i,v in enumerate(row)])
    ws.auto_filter.ref=f'A4:E{4+len(evidence)}';counts['Evidence Files']=len(evidence)
    temporary=parent/'.~summary.xlsx';destination=parent/'summary.xlsx'
    book.save(temporary)
    check=load_workbook(temporary,read_only=True,data_only=True)
    assert check.sheetnames[0]=='Overview' and len(list(check['Scores by Split'].values))==12
    check.close();temporary.replace(destination)
    return {'workbook':str(destination.relative_to(root)),'sheets':len(book.sheetnames),'table_rows':sum(counts.values())}


if __name__=='__main__':
    print(json.dumps(build_summary(),indent=2))

"""Build one readable MLP workbook for baseline and tuning; never fit or load binary models.

Layout: multilayer_perceptron/summary.xlsx sits beside baseline/ and tuning/, which
hold each variant's saved files directly. The display matches the random forest,
XGBoost, SVR and dummy summaries, plus MLP-only sheets (seeds, epochs, histories).
"""
from pathlib import Path
from collections import defaultdict
from statistics import mean, stdev
import csv
import hashlib
import json
import math
import os
import re
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter
from fetch_models import require_models

ROOT = Path(__file__).resolve().parents[1]
MLP = Path('provenance/models/multilayer_perceptron')
VARIANTS = ['baseline', 'tuning']
# Saved CSV -> workbook sheet. Files with the same sheet name are combined, and a
# 'variant' column records which folder (baseline/ or tuning/) each row came from.
TABLES = {
    'metrics.csv': 'Scores by Split',
    'seed_summary.csv': 'Score Summary',
    'baseline_comparison.csv': 'Baseline Comparison',
    'candidate_summary.csv': 'Search Ranking',
    'candidate_scores.csv': 'Search by Seed',
    'validation_subgroups.csv': 'Saved Similarity',
    'predictions.csv': 'Predictions',
    'subgroup_summary.csv': 'Similarity Summary',
    'cliff_summary.csv': 'Cliff Summary',
    'cliff_molecule_group_summary.csv': 'Cliff Molecule Summary',
    'cliff_results.csv': 'Cliff Results by Seed',
    'cliff_molecule_groups.csv': 'Cliff Molecules by Seed',
    'cliff_pairs.csv': 'Cliff Pairs',
    'epoch_selections.csv': 'Selected Epochs',
    'internal_history.csv': 'Stopping History',
    'selected_internal_history.csv': 'Stopping History',
    'full_refit_history.csv': 'Refit History',
    'candidate_history.csv': 'Search History',
    'internal_assignments.csv': 'Internal Membership',
}
SUBTITLES = {
    'Scores by Split': 'Exact saved scores. Each row is one initialization seed on one split and subset.',
    'Score Summary': 'Means and sample standard deviations across the five initialization seeds.',
    'Baseline Comparison': 'Baseline and tuned scores on the same frozen data. Compare matching seeds and splits.',
    'Search Ranking': 'Every candidate, summarized across search seeds on the internal stopping set (training data only).',
    'Search by Seed': 'Every candidate fit and its internal stopping score. These are not outer-validation scores.',
    'Saved Similarity': 'Validation errors by similarity to the nearest training molecule, for each seed.',
    'Predictions': 'Every saved training and validation prediction; one row per molecule, split and seed.',
    'Similarity Summary': 'Saved similarity-group errors, summarized across seeds.',
    'Cliff Summary': 'Saved performance on similar molecule pairs with large measured affinity differences, across seeds.',
    'Cliff Molecule Summary': 'Saved validation error on cliff-involved versus other molecules, across seeds.',
    'Cliff Results by Seed': 'Cliff-pair errors and direction accuracy for each seed.',
    'Cliff Molecules by Seed': 'Cliff-involved versus other molecules, for each seed.',
    'Cliff Pairs': 'Frozen validation pairs defining the activity-cliff diagnostic (reused by tuning).',
    'Selected Epochs': 'Training length chosen on the internal stopping set, before the full-training refit.',
    'Stopping History': 'Epoch-by-epoch internal fitting and stopping losses (pKi²).',
    'Refit History': 'Epoch history for fresh models fitted to all outer-training molecules.',
    'Search History': 'Epoch histories for every searched candidate and seed.',
    'Internal Membership': 'Outer-training molecules assigned to internal fitting or stopping sets (reused by tuning).',
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text())


def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8') as stream:
        return list(csv.DictReader(stream))


def flatten(value, key=''):
    if isinstance(value, dict):
        for name, item in value.items():
            yield from flatten(item, f'{key}.{name}' if key else name)
    elif isinstance(value, list):
        # A short list is easier to read together than across many technical rows.
        if all(not isinstance(item, (dict, list)) for item in value):
            yield key, json.dumps(value)
        else:
            for i, item in enumerate(value):
                yield from flatten(item, f'{key}[{i}]')
    else:
        yield key, value


def numeric(value):
    # Preserve long identifiers, molecular strings and formula-like text literally.
    if value == '':
        return None
    if re.fullmatch(r'-?\d{1,14}', value):
        return int(value)
    if re.fullmatch(r'-?(?:\d+\.\d*|\d*\.\d+)(?:[eE][+-]?\d+)?', value):
        number = float(value)
        if math.isfinite(number):
            return number
    return value


def build_summary(root=ROOT):
    root = Path(root); parent = root/MLP
    require_models(root, MLP)  # Fitted models live outside git; fetch before hashing.
    # Tuning is absent right after a new baseline is accepted (it was derived from the
    # previous baseline). The workbook then shows the baseline and marks tuning pending.
    variants = [v for v in VARIANTS if (parent/v/'manifest.json').is_file()]
    assert 'baseline' in variants, 'Run notebook 08 first: no accepted MLP baseline exists.'
    manifests = {v: read_json(parent/v/'manifest.json') for v in variants}
    baseline, tuning = manifests['baseline'], manifests.get('tuning')
    for m in manifests.values():
        assert m.get('completed_at_utc') and not m.get('test_evaluation_performed')
        assert m['source_preparation_manifest'] == baseline['source_preparation_manifest']
        assert digest(root/m['source_preparation_manifest']) == m['source_preparation_manifest_sha256']
        for field in ['output_sha256', 'source_artifact_sha256', 'reused_baseline_artifact_sha256']:
            for name, expected in m.get(field, {}).items():
                assert digest(root/name) == expected, f'Changed evidence: {name}'
    if tuning:
        assert tuning['baseline_manifest'] == (MLP/'baseline/manifest.json').as_posix()
        assert digest(root/tuning['baseline_manifest']) == tuning['baseline_manifest_sha256']
    prep_path = root/baseline['source_preparation_manifest']; prep = read_json(prep_path)
    config = baseline['configuration']; params = config['model_parameters']
    metrics = {v: read_csv(parent/v/'metrics.csv') for v in variants}
    # Recompute the seed summary independently and confirm it matches the saved one.
    saved_summary = {(r['split_strategy'], r['model_variant'], r['subset']): r for v in variants for r in read_csv(parent/v/'seed_summary.csv')}
    book = Workbook(write_only=True); counts = {}

    def cell(ws, value, style=None, link=None):
        c = WriteOnlyCell(ws, value=value)
        if isinstance(value, str): c.data_type = 's'
        c.font = Font(name='Calibri', size=19 if style == 'title' else 11, bold=bool(style), color='FFFFFF' if style else '18354A')
        c.alignment = Alignment(vertical='top', wrap_text=True)
        if style: c.fill = PatternFill('solid', fgColor='18354A' if style == 'title' else '007F82')
        if isinstance(value, float): c.number_format = '0.0000'
        if link:
            c.hyperlink = Hyperlink(ref=c.coordinate, location=link[1:]) if link.startswith('#') else link
            c.font = Font(name='Calibri', size=11, color='007F82', underline='single')
        return c

    def worksheet(name, widths):
        ws = book.create_sheet(name); ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = '007F82'
        for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    def table(name, subtitle, headers, rows, widths=None):
        ws = worksheet(name, widths or [max(19, min(58, len(h)+3)) for h in headers]); ws.freeze_panes = 'A5'
        for row in [1, 2]: ws.merged_cells.add(f'A{row}:{get_column_letter(len(headers))}{row}')
        ws.row_dimensions[1].height = 32; ws.row_dimensions[2].height = 42; ws.row_dimensions[4].height = 40
        ws.append([cell(ws, name, 'title')]); ws.append([cell(ws, subtitle)])
        ws.append([cell(ws, 'Back to overview', link="#'Overview'!A1")]); ws.append([cell(ws, h, 'header') for h in headers])
        count = 0
        for i, row in enumerate(rows, 5):
            if name in ['Training Guide', 'Settings']: ws.row_dimensions[i].height = 44
            ws.append([cell(ws, v) for v in row]); count += 1
        ws.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{4+count}'; counts[name] = count
        return ws

    # Combine each saved CSV across both folders, keeping every original column.
    grouped = {}
    for variant in variants:
        for path in sorted((parent/variant).glob('*.csv')):
            assert path.name in TABLES, f'Add a readable sheet name for {path.name}'
            name = TABLES[path.name]; rows = read_csv(path)
            bucket = grouped.setdefault(name, {'headers': ['variant'], 'rows': []})
            for key in rows[0] if rows else []:
                if key not in bucket['headers']: bucket['headers'].append(key)
            bucket['rows'].extend({'variant': variant, **r} for r in rows)

    cover = worksheet('Overview', [19, 20, 18, 17, 19, 19, 19, 30]); cover.freeze_panes = 'D6'
    def line(values, style=None, height=30, links=None):
        line.row += 1; cover.row_dimensions[line.row].height = height
        if len(values) < 8:
            first = 'A' if len(values) == 1 else 'B' if len(values) == 2 else 'C'
            cover.merged_cells.add(f'{first}{line.row}:H{line.row}')
        cover.append([cell(cover, v, style, (links or {}).get(i)) for i, v in enumerate(values)])
    line.row = 0
    line(['MLP | BASELINE + TUNED'], 'title', 38)
    tuned_time = tuning['completed_at_utc'] if tuning else 'pending (run notebook 09)'
    line(['Saved executions', f"Baseline: {baseline['completed_at_utc']}  |  Tuned: {tuned_time}"], height=36)
    seeds = config['initialization_seeds']
    line([f"{len(seeds)} initialization seeds per split and variant. Scores are seed means; the last column shows the sample SD (seed variability, not a confidence interval)."], height=36)
    line(['Validation estimates generalization; training describes fit. Reserved test sets remain unevaluated.'])
    line(['Split', 'Subset', 'Variant', 'Molecules', 'MAE (pKi)', 'RMSE (pKi)', 'R²', 'Seed SD (MAE / RMSE / R²)'], 'header')
    for subset in ['validation', 'train']:
        for split in ['random', 'scaffold']:
            for variant in variants:
                rows = [r for r in metrics[variant] if r['split_strategy'] == split and r['subset'] == subset]
                assert len(rows) == len(seeds)
                means = {k: mean(float(r[k]) for r in rows) for k in ['mae_pki', 'rmse_pki', 'r2']}
                sds = {k: stdev(float(r[k]) for r in rows) for k in means}
                saved = saved_summary[split, rows[0]['model_variant'], subset]
                for k in means:
                    assert math.isclose(means[k], float(saved[k+'_mean']), abs_tol=1e-12)
                    assert math.isclose(sds[k], float(saved[k+'_sample_sd']), abs_tol=1e-12)
                line([split.title(), subset.title(), 'Tuned' if variant == 'tuning' else 'Baseline', int(rows[0]['n_structures']),
                      means['mae_pki'], means['rmse_pki'], means['r2'],
                      f"{sds['mae_pki']:.4f} / {sds['rmse_pki']:.4f} / {sds['r2']:.4f}"])
    line(['Lower MAE / RMSE is better; higher R² is better. A tuned model is not necessarily better on validation.'])
    line(['MODEL AND TRAINING AT A GLANCE'], 'header')
    fp = prep['configuration']['fingerprint']; layers = params['hidden_layer_sizes']
    line(['Architecture', 'Neural network (MLP)', f"{config['input_bits']} inputs → hidden layers {layers} → 1 output. ReLU hidden units; linear pKi output."], height=42)
    line(['Inputs', 'Morgan fingerprints', f"{fp['bits']} bits; radius {fp['radius']}; chirality {fp['include_chirality']}"], height=38)
    line(['Baseline', 'Fixed before fitting', f"hidden layers = {layers}; alpha (L2) = {params['alpha']}; learning rate = {params['learning_rate_init']}; batch = {params['batch_size']}"], height=42)
    if tuning:
        for split, selected in tuning['selection'].items():
            chosen = {k: selected[k] for k in ['hidden_layers', 'alpha', 'learning_rate_init']}
            line([split.title()+' tuned', 'Selected parameters', json.dumps(chosen)], height=42)
            line([split.title()+' search', 'Mean stopping MSE', f"{selected['mean_best_stopping_mse_pki2']:.4f} pKi² (internal stopping set; training data only)"], height=36)
        search = tuning['configuration']
        line(['Tuning', f"{len(search['grid'])} candidates / split", f"Each candidate trained with seeds {', '.join(map(str, search['search_seeds']))}; lowest mean internal-stopping MSE wins. Outer validation is not used."], height=46)
    else:
        line(['Tuning', 'Pending', 'The previous tuning was derived from an older baseline and was cleared. Run notebook 09.'], height=42)
    line(['Training', 'Adam + internal stopping', f"Epochs chosen on an internal stopping set carved from training (at most {config['maximum_epochs']}; patience {config['patience']}); then a fresh model is refit on all training molecules."], height=46)
    line(['Dataset', f"{prep['counts']['structures']:,} structures", f"{prep['counts']['retained_measurements']:,} measurements; frozen random and scaffold splits."], height=36)
    line(['Scope', 'Development validation', 'Prior development inspected validation labels; final claims still require reserved test evaluation.'], height=44)
    line(['READ NEXT'], 'header')
    navigation = [('Training Guide', 'Architecture, parameter meanings, data sizes, and training procedure.'),
                  ('Search Ranking', 'Every candidate and its internal stopping score (training data only).'),
                  ('Baseline Comparison', 'Exact saved baseline-versus-tuned results, seed by seed.'),
                  ('Similarity Summary', 'Errors by nearest-training similarity, across seeds.'),
                  ('Cliff Summary', 'Errors on frozen validation activity-cliff pairs, across seeds.'),
                  ('Selected Epochs', 'How long each network trained, chosen on the internal stopping set.'),
                  ('Evidence Files', 'Original files (including loss and prediction figures), checksums, and sheet locations.')]
    for name, text in navigation:
        if name in grouped or name in ['Training Guide', 'Evidence Files']:
            line([name, text], links={0: f"#'{name}'!A1"}, height=34)
    cover.sheet_properties.pageSetUpPr.fitToPage = True; cover.page_setup.orientation = 'landscape'; cover.page_setup.paperSize = '8'
    cover.page_setup.fitToWidth = 1; cover.page_setup.fitToHeight = 1; cover.print_area = f'A1:H{line.row}'

    guide = [['Both', 'Algorithm', 'Multilayer perceptron (MLPRegressor)', 'A fully connected neural network: each hidden layer combines the previous layer\'s values; the output is one continuous pKi.'],
             ['Both', 'Target', 'pKi', prep['target_policy']+'; higher pKi means stronger affinity. One pKi unit is a tenfold change in Ki.'],
             ['Both', 'Feature encoding', 'Binary Morgan fingerprints', json.dumps(fp)],
             ['Both', 'Optimizer', 'Adam', 'Minimizes squared prediction error with an L2 penalty (alpha) on the weights.'],
             ['Both', 'Activation', 'ReLU', 'Hidden units output max(0, x); the regression output is linear.'],
             ['Both', 'Internal stopping', 'Fit and stopping sets from training only', 'Outer-training molecules are split again; outer validation is held aside.'],
             ['Both', 'Epoch selection', config['epoch_selection'], f"At most {config['maximum_epochs']} epochs; patience {config['patience']}."],
             ['Both', 'Final training', config['full_refit'], 'The selected epoch count is reused with a fresh model on all outer-training molecules.'],
             ['Both', 'Initialization seeds', ', '.join(map(str, seeds)), 'Independent random starts on the same data, not different splits or experiments. Their SD shows seed sensitivity.'],
             ['Baseline', 'Fit', 'Fixed settings', 'One network per seed and outer split, with the settings below.']]
    if tuning:
        search = tuning['configuration']
        guide += [['Tuned', 'Search', 'Fixed candidate grid', search['selection_rule']],
                  ['Tuned', 'Search seeds', ', '.join(map(str, search['search_seeds'])), 'Seeds used to score each candidate on the internal stopping set.'],
                  ['Tuned', 'Search grid', f"{len(search['grid'])} candidates", json.dumps([{k: g[k] for k in ['hidden_layers', 'alpha', 'learning_rate_init']} for g in search['grid']])],
                  ['Tuned', 'Reused from baseline', 'Internal membership and cliff pairs', 'Tuning reuses the baseline\'s frozen internal split and cliff pairs, so both variants are scored on identical data.'],
                  ['Tuned', 'Final fit', 'Selected settings, all final seeds', 'Each split\'s chosen configuration is refit with every initialization seed.']]
    guide += [['Both', 'MAE', 'Mean absolute error', 'Typical absolute prediction error in pKi; lower is better.'],
              ['Both', 'RMSE', 'Root mean squared error', 'Emphasizes large errors; lower is better.'],
              ['Both', 'R²', 'Coefficient of determination', '1 is perfect; zero matches the evaluated-subset mean; negative values are possible.'],
              ['Both', 'Stopping MSE', 'pKi² units', 'Mean squared error on the internal stopping set, used only to choose epochs and candidates.'],
              ['Both', 'Cliffs', 'Saved diagnostics', 'Unordered validation pairs with Tanimoto ≥ 0.8 and absolute observed pKi gap ≥ 1: '+json.dumps(config['cliff_definition'])],
              ['Both', 'Similarity', 'Saved diagnostics', 'Validation errors grouped by exact fingerprint match and nearest-training Tanimoto (0.80 cutoff).'],
              ['Both', 'Figures', 'PNG files', 'Loss curves and prediction plots for the representative seed are linked from Evidence Files.'],
              ['Both', 'Models', 'Compressed binaries', 'models.joblib.gz files are preserved and linked; this report does not load or retrain them.']]
    explanations = {'hidden_layer_sizes': 'Units in each hidden layer, from input side to output side.',
                    'alpha': 'L2 penalty on weights; larger values give smoother, more conservative networks.',
                    'learning_rate_init': 'Adam step size; smaller values learn more slowly but steadily.',
                    'batch_size': 'Molecules per weight update.', 'activation': 'Hidden-unit function (relu).',
                    'solver': 'Optimizer (adam).', 'loss': 'Training loss (squared error).',
                    'shuffle': 'Shuffle molecules each epoch.',
                    'early_stopping': 'scikit-learn\'s built-in stopping is off; the notebook\'s internal stopping set is used instead.',
                    'max_iter': 'Epochs per call; the notebook trains one epoch at a time to record histories.',
                    'n_iter_no_change': 'Set above the epoch limit so scikit-learn never stops training on its own.'}
    for variant in variants:
        for split in ['random', 'scaffold']:
            values = dict(params)
            if variant == 'tuning':
                selected = tuning['selection'][split]
                values.update(hidden_layer_sizes=selected['hidden_layers'], alpha=selected['alpha'], learning_rate_init=selected['learning_rate_init'])
            for key, value in values.items():
                guide.append([variant, split+' / '+key, json.dumps(value), explanations.get(key, 'Saved scikit-learn estimator setting.')])
    for split, subsets in prep['counts']['split_subsets'].items():
        for subset, sizes in subsets.items(): guide.append(['Both', 'Dataset sizes', split+' / '+subset, sizes['structures']])
    for split, info in config['internal_split'].items():
        for key, value in info.items(): guide.append(['Both', 'Internal split', split+' / '+key, value])
    table('Training Guide', 'Plain-English descriptions and complete fitted-model parameters.', ['Variant', 'Topic', 'Choice / value', 'Explanation'], guide, [18, 38, 62, 100])
    settings = []
    for variant in variants:
        for path in sorted((parent/variant).glob('*.json')):
            settings.extend([f'{variant}/{path.name}', key, value] for key, value in flatten(read_json(path)))
    settings.extend(['preparation/manifest.json', key, value] for key, value in flatten(prep))
    table('Settings', 'Saved JSON values, software versions and verification records.', ['Source', 'Setting', 'Value'], settings, [38, 75, 110])
    # Same order as the other model summaries first, then MLP-only sheets.
    order = ['Scores by Split', 'Score Summary', 'Baseline Comparison', 'Search Ranking', 'Search by Seed', 'Saved Similarity', 'Predictions']
    for name in order:
        if name not in grouped: continue
        data = grouped[name]; headers = data['headers']
        table(name, SUBTITLES[name]+' Variant identifies the producing folder.', headers,
              [[numeric(str(r.get(h, ''))) for h in headers] for r in data['rows']])
    if tuning:
        selected = [[split, s['candidate_index'], json.dumps(s['hidden_layers']), s['alpha'], s['learning_rate_init'], s['trainable_parameters'], s['mean_best_stopping_mse_pki2']]
                    for split, s in tuning['selection'].items()]
        table('Selected Parameters', 'Selections from the saved tuning manifest. Stopping MSE uses internal training data only.',
              ['Split', 'Candidate', 'Hidden layers', 'Alpha (L2)', 'Learning rate', 'Trainable parameters', 'Mean stopping MSE (pKi²)'], selected, [18, 16, 22, 18, 18, 22, 26])
    for name in dict.fromkeys(TABLES.values()):
        if name in order or name not in grouped: continue
        data = grouped[name]; headers = data['headers']
        table(name, SUBTITLES[name]+' Variant identifies the producing folder.', headers,
              [[numeric(str(r.get(h, ''))) for h in headers] for r in data['rows']])
    # Evidence is linked relative to the workbook, keeping links portable with this folder.
    evidence = [p for v in variants for p in sorted((parent/v).iterdir()) if p.is_file()] + [prep_path]
    ws = table('Evidence Files', 'Every backing file, plus the preparation source. Binaries and figures are linked, not expanded.',
               ['File', 'Purpose', 'Workbook sheet', 'Bytes', 'SHA-256'], [], [52, 52, 30, 16, 72])
    for path in evidence:
        tab = TABLES.get(path.name, 'Settings' if path.suffix == '.json' else 'Linked file only')
        target = Path(os.path.relpath(path, parent)).as_posix()
        purpose = ('Fitted models; compressed binary' if path.name.endswith('.gz') else 'Diagnostic figure' if path.suffix == '.png'
                   else 'Preparation source' if path == prep_path else 'Original saved evidence')
        row = [target, purpose, tab, path.stat().st_size, digest(path)]
        ws.append([cell(ws, v, link=target if i == 0 else f"#'{v}'!A1" if i == 2 and v != 'Linked file only' else None) for i, v in enumerate(row)])
    ws.auto_filter.ref = f'A4:E{4+len(evidence)}'; counts['Evidence Files'] = len(evidence)
    temporary = parent/'.~summary.xlsx'; destination = parent/'summary.xlsx'
    book.save(temporary)
    # Reopen the finished file before replacing a reader's existing summary.
    check = load_workbook(temporary, read_only=True, data_only=True)
    assert check.sheetnames[0] == 'Overview' and len(list(check['Scores by Split'].values)) == 4 + sum(len(m) for m in metrics.values())
    check.close(); temporary.replace(destination)
    return {'workbook': destination.relative_to(root).as_posix(), 'sheets': len(book.sheetnames), 'table_rows': sum(counts.values())}


if __name__ == '__main__':
    print(json.dumps(build_summary(), indent=2))

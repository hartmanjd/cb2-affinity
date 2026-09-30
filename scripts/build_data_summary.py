"""Build the CB2 data provenance workbook from the current pipeline files; never edit data.

The workbook shows where every number in the dataset came from: the ChEMBL download,
each filtering step, every kept and removed measurement with its reason, and the
assays and documents (papers and patents) behind them. Every table is generated
from the saved raw, selection, curation and preparation records, which are
checksum-verified first, so the workbook always matches the current data.

The file is written deterministically (fixed timestamps), so rebuilding it from
unchanged data gives identical bytes and does not show up as a change in git.
"""
from pathlib import Path
from collections import Counter
import io
import json
import os
import re
import zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter
from build_mlp_summary import digest, read_json, flatten
from provenance_support import current_curation, current_preparation

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path('provenance/CB2_data_provenance.xlsx')
MODEL_SUMMARIES = [('Dummy baselines', 'dummy_baselines'), ('Random forest', 'random_forest'),
                   ('XGBoost', 'xgboost'), ('Support vector regression', 'support_vector_regression'),
                   ('Multilayer perceptron', 'multilayer_perceptron')]
# Plain-English meaning of each funnel step, keyed by the saved stage name.
STEP_MEANING = {
    'Downloaded target activities': 'Every activity record ChEMBL links to human CB2 (CHEMBL253), with no filtering.',
    'Ki endpoint': 'Keeps binding-affinity Ki measurements; drops IC50, EC50, efficacy and other endpoints.',
    'Binding assay type B': 'Keeps binding assays; functional and other assay types measure something else.',
    'Direct target confidence 9': 'Keeps assays that ChEMBL maps directly to the single CB2 protein.',
    'Human assay organism': 'Keeps experiments run on the human receptor.',
    'Essential fields and valid nM Ki': 'Drops rows missing a structure, value or nM unit, or with an invalid Ki.',
    'Exact Ki measurements': 'Drops censored values such as ">10000 nM", which are limits rather than measurements.',
    'Acceptable validity status': 'Drops rows ChEMBL flags as potentially invalid.',
    'No potential-duplicate flag': 'Drops rows ChEMBL flags as possible duplicates of another record.',
    'Unique activity records': 'Confirms no activity record appears twice.',
    'After structure review': 'Sets aside multi-component structures (e.g. salts with several parts) for review.',
    'After conflicting measurements review': 'Sets aside structures whose measurements disagree by 1 pKi unit or more.',
    'Eligible structures before conflict review': 'Distinct canonical structures before removing conflicting ones.',
    'Final molecules (canonical structure keys)': 'Final molecules used for modeling; each has one aggregated pKi.',
}
# Curation records reasons as short codes; show readers what they mean.
REASON_TEXT = {
    'pki_range_ge_1.0': "Measurements of this structure disagree by 1 pKi unit or more (pki_range_ge_1.0)",
    'multiple_components_review': "Structure has several disconnected parts, e.g. a salt (multiple_components_review)",
}


def verified_sources(root):
    """Return the current stage manifests after checking their recorded checksums."""
    curation_path = root/current_curation(root)
    preparation_path = root/current_preparation(root)
    curation, preparation = read_json(curation_path), read_json(preparation_path)
    # Every upstream file the curation and preparation recorded must be unchanged.
    for record in [curation['file_sha256'], preparation['source_sha256'], preparation['artifact_sha256']]:
        for name, expected in record.items():
            assert digest(root/name) == expected, f'Changed source: {name}'
    selection_path = root/curation['selection_manifest']
    raw_folder = root/curation['source_raw_folder']
    return {'curation': curation_path, 'preparation': preparation_path,
            'selection': selection_path, 'acquisition': raw_folder/'acquisition_manifest.json',
            'raw_folder': raw_folder}


def normalize_xlsx(data, stamp):
    """Rewrite an .xlsx so identical content always gives identical bytes.

    Excel files are zip archives; openpyxl records the save time both in the zip
    entries and in docProps/core.xml. Both are replaced with fixed values.
    """
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == 'docProps/core.xml':
                content = re.sub(rb'(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)',
                                 rb'\g<1>' + stamp.encode() + rb'\g<3>', content)
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(fixed, content)
    return output.getvalue()


def render(root=ROOT):
    """Build the workbook in memory and return its normalized bytes."""
    from build_project_reference import build_measurement_tables
    root = Path(root)
    paths = verified_sources(root)
    acquisition, selection = read_json(paths['acquisition']), read_json(paths['selection'])
    curation, preparation = read_json(paths['curation']), read_json(paths['preparation'])
    status = read_json(paths['raw_folder']/'chembl_status.json')
    # The reference builder reconciles every downloaded activity to a kept or
    # removed decision and re-derives each molecule's final pKi; reuse it here.
    activities, assays, documents, kept, removed, funnel = build_measurement_tables(
        {'curation': paths['curation'].relative_to(root).as_posix()})
    parent = root/WORKBOOK.parent
    book = Workbook(write_only=True); counts = {}

    def cell(ws, value, style=None, link=None):
        c = WriteOnlyCell(ws, value=value)
        if isinstance(value, str): c.data_type = 's'
        c.font = Font(name='Calibri', size=19 if style == 'title' else 11, bold=bool(style), color='FFFFFF' if style else '18354A')
        c.alignment = Alignment(vertical='top', wrap_text=True)
        if style: c.fill = PatternFill('solid', fgColor='18354A' if style == 'title' else '007F82')
        if isinstance(value, float): c.number_format = '0.0000'
        if link:
            c.hyperlink = Hyperlink(ref=c.coordinate, location=link[1:]) if link.startswith('#') else link
            c.font = Font(name='Calibri', size=11, color='007F82', underline='single')
        return c

    def worksheet(name, widths):
        ws = book.create_sheet(name); ws.sheet_view.showGridLines = False; ws.sheet_properties.tabColor = '007F82'
        for i, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    def table(name, subtitle, headers, rows, widths=None, links=None, counts_columns=()):
        ws = worksheet(name, widths or [max(16, min(50, len(h)+3)) for h in headers]); ws.freeze_panes = 'A5'
        for row in [1, 2]: ws.merged_cells.add(f'A{row}:{get_column_letter(len(headers))}{row}')
        ws.row_dimensions[1].height = 32; ws.row_dimensions[2].height = 42; ws.row_dimensions[4].height = 40
        ws.append([cell(ws, name, 'title')]); ws.append([cell(ws, subtitle)])
        ws.append([cell(ws, 'Back to overview', link="#'Overview'!A1")]); ws.append([cell(ws, h, 'header') for h in headers])
        count = 0
        for row in rows:
            cells = [cell(ws, v, link=(links(row, i) if links else None)) for i, v in enumerate(row)]
            for i in counts_columns:
                cells[i].number_format = '#,##0'  # Thousands separators for counts, not IDs.
            ws.append(cells); count += 1
        ws.auto_filter.ref = f'A4:{get_column_letter(len(headers))}{4+count}'; counts[name] = count
        return ws

    # ---- Overview -------------------------------------------------------------
    cover = worksheet('Overview', [30, 24, 22, 70]); cover.freeze_panes = 'A3'
    def line(values, style=None, height=28, links=None):
        line.row += 1; cover.row_dimensions[line.row].height = height
        if len(values) < 4:
            first = 'A' if len(values) == 1 else 'B' if len(values) == 2 else 'C'
            cover.merged_cells.add(f'{first}{line.row}:D{line.row}')
        cover.append([cell(cover, v, style, (links or {}).get(i)) for i, v in enumerate(values)])
    line.row = 0
    line(['CB2 DATA PROVENANCE'], 'title', 38)
    line(['Where every number in the dataset came from. Generated from the current pipeline files by scripts/build_data_summary.py; every source is checksum-verified.'], height=36)
    line(['KEY FACTS'], 'header')
    prep_counts = preparation['counts']
    facts = [('Target', f"{acquisition['target_name']} ({acquisition['target_chembl_id']}), {acquisition['organism']}, UniProt {acquisition['uniprot']}"),
             ('Source database', f"{status['chembl_db_version']}, released {status['chembl_release_date']} (ChEMBL REST API)"),
             ('Downloaded (UTC)', acquisition['retrieved_at_utc']),
             ('Downloaded records', f"{acquisition['activity_records']:,} activities, {acquisition['assay_records']:,} assays, {acquisition['document_records']:,} documents"),
             ('Kept measurements', f"{len(kept):,} Ki measurements from {len({r['assay_chembl_id'] for r in kept}):,} assays and {len({r['document_chembl_id'] for r in kept}):,} documents"),
             ('Final molecules', f"{prep_counts['structures']:,} structures, each with one aggregated pKi target"),
             ('Target value', curation['policies']['target_aggregation']),
             ('Produced on', f"{curation['software_versions'].get('platform', 'unknown OS')}, Python {curation['software_versions']['python']}")]
    for label, value in facts:
        line([label, value], height=30)
    line(['HOW THE DATA WAS NARROWED'], 'header')
    line(['Step', 'Remaining', 'Removed here', 'Meaning'], 'header')
    for step in funnel:
        removed_here = f"{step['removed_at_step']:,}" if step['removed_at_step'] else ''
        line([step['stage'], f"{step['remaining']:,} {step['unit']}", removed_here, STEP_MEANING.get(step['stage'], '')], height=32)
    line(['READ NEXT'], 'header')
    for name, text in [('Filter Funnel', 'Every step with its count and the file that records it.'),
                       ('Removal Reasons', 'How many measurements were removed for each reason.'),
                       ('Kept Measurements', 'All kept measurements with their paper, assay and final molecule pKi.'),
                       ('Removed Measurements', 'Every removed measurement and exactly why.'),
                       ('Assays', 'Every downloaded experiment and whether the final dataset uses it.'),
                       ('Documents', 'Every paper and patent ChEMBL curated these measurements from.'),
                       ('Pipeline Settings', 'Stage completion records, policies, fingerprints and split settings.'),
                       ('Evidence Files', 'Source files with sizes and SHA-256 checksums.')]:
        line([name, text], links={0: f"#'{name}'!A1"}, height=28)
    line(['MODEL RESULTS'], 'header')
    for label, folder in MODEL_SUMMARIES:
        target_file = f'models/{folder}/summary.xlsx'
        line([label, 'Open the model summary workbook'], links={0: target_file}, height=26)
    cover.print_area = f'A1:D{line.row}'

    # ---- Detail sheets -------------------------------------------------------------
    table('Filter Funnel', 'Each step, its remaining count, and the saved record that documents it. Molecule counts start after measurement filtering.',
          ['Step', 'Unit', 'Remaining', 'Removed at this step', 'Meaning', 'Recorded in'],
          [[s['stage'], s['unit'], s['remaining'], s['removed_at_step'] or None, STEP_MEANING.get(s['stage'], ''), s['source']] for s in funnel],
          [42, 14, 14, 18, 70, 80], counts_columns=(2, 3))
    for r in removed:
        r['exclusion_reason'] = REASON_TEXT.get(r['exclusion_reason'], r['exclusion_reason'])
    reasons = Counter(r['exclusion_reason'] for r in removed)
    table('Removal Reasons', f"All {len(removed):,} removed measurements grouped by reason; together with {len(kept):,} kept they account for every one of the {len(activities):,} downloaded activities.",
          ['Reason', 'Measurements'], [[reason, n] for reason, n in reasons.most_common()], [90, 18], counts_columns=(1,))
    kept_columns = ['activity_id', 'molecule_chembl_id', 'rdkit_smiles', 'standard_value', 'pki', 'final_molecule_pki_target',
                    'final_molecule_measurement_count', 'assay_chembl_id', 'assay_description', 'document_chembl_id', 'year', 'paper_title', 'doi', 'pubmed_id']
    kept_headers = ['Activity ID', 'Molecule ChEMBL ID', 'Standardized structure (RDKit SMILES)', 'Ki (nM)', 'pKi (this measurement)', 'Final pKi for this molecule',
                    'Measurements behind final pKi', 'Assay ChEMBL ID', 'Assay description', 'Document ChEMBL ID', 'Year', 'Paper title', 'DOI', 'PubMed ID']
    table('Kept Measurements', 'Every measurement in the final dataset. pKi = 9 - log10(Ki in nM). A molecule\'s final pKi is the median within each assay, then the median across assays.',
          kept_headers, [[r.get(k) if k != 'standard_value' else float(r[k]) for k in kept_columns] for r in sorted(kept, key=lambda r: r['activity_id'])],
          [14, 18, 50, 12, 14, 14, 14, 18, 60, 18, 8, 60, 28, 12],
          links=lambda row, i: f"https://doi.org/{row[12]}" if i == 12 and row[12] else None)
    removed_columns = ['activity_id', 'molecule_chembl_id', 'standard_type', 'standard_relation', 'standard_value', 'standard_units',
                       'assay_chembl_id', 'document_chembl_id', 'year', 'exclusion_reason', 'decision_source']
    table('Removed Measurements', 'Every downloaded measurement that is not in the final dataset, with the reason and the saved record that removed it.',
          ['Activity ID', 'Molecule ChEMBL ID', 'Measurement type', 'Relation', 'Value', 'Units', 'Assay ChEMBL ID', 'Document ChEMBL ID', 'Year', 'Reason', 'Recorded in'],
          [[r.get(k) for k in removed_columns] for r in sorted(removed, key=lambda r: r['activity_id'])],
          [14, 18, 16, 10, 12, 10, 18, 18, 8, 70, 70])
    assay_columns = ['assay_chembl_id', 'assay_type', 'assay_organism', 'confidence_score', 'confidence_description', 'relationship_description',
                     'bao_label', 'assay_cell_type', 'description', 'document_chembl_id', 'downloaded_measurements', 'kept_measurements', 'used_in_final_dataset']
    table('Assays', 'Every downloaded CB2 experiment. Assay type B = binding; confidence 9 = mapped directly to the single CB2 protein.',
          ['Assay ChEMBL ID', 'Type', 'Organism', 'Confidence', 'Confidence meaning', 'Relationship', 'Format', 'Cell type', 'Description', 'Document',
           'Downloaded measurements', 'Kept measurements', 'Used in final dataset'],
          [[r.get(k) for k in assay_columns] for r in sorted(assays, key=lambda r: (-r['kept_measurements'], r['assay_chembl_id']))],
          [18, 8, 16, 11, 36, 28, 22, 16, 70, 18, 14, 14, 12])
    document_columns = ['document_chembl_id', 'doc_type', 'year', 'journal', 'title', 'authors', 'doi', 'pubmed_id', 'patent_id',
                        'downloaded_measurements', 'kept_measurements', 'used_in_final_dataset']
    table('Documents', 'Every paper, patent and dataset behind the downloaded measurements. ChEMBL curators extracted the measurements from these sources; DOIs link to the originals.',
          ['Document ChEMBL ID', 'Type', 'Year', 'Journal', 'Title', 'Authors', 'DOI', 'PubMed ID', 'Patent ID',
           'Downloaded measurements', 'Kept measurements', 'Used in final dataset'],
          [[r.get(k) for k in document_columns] for r in sorted(documents, key=lambda r: (-r['kept_measurements'], r['document_chembl_id']))],
          [18, 14, 8, 22, 70, 50, 28, 12, 18, 14, 14, 12],
          links=lambda row, i: f"https://doi.org/{row[6]}" if i == 6 and row[6] else None)
    settings = []
    for label, path in [('acquisition', paths['acquisition']), ('selection', paths['selection']),
                        ('curation', paths['curation']), ('preparation', paths['preparation'])]:
        for key, value in flatten(read_json(path)):
            # Long checksum and query lists are in Evidence Files; keep this sheet readable.
            if key.startswith(('file_sha256', 'source_sha256', 'artifact_sha256', 'document_queries')):
                continue
            settings.append([label, key, value if not isinstance(value, (dict, list)) else json.dumps(value)])
    table('Pipeline Settings', 'Completion records, policies, fingerprint and split settings, and software for each data stage (notebooks 00-03).',
          ['Stage', 'Setting', 'Value'], settings, [16, 70, 90])
    evidence = [paths['acquisition'], *(paths['raw_folder']/n for n in ['chembl_status.json', 'chembl_cb2_target.json', 'chembl_cb2_activity.json',
                                                                    'chembl_cb2_assay.json', 'chembl_cb2_documents.json']),
                paths['selection'], paths['selection'].parent/selection['selected_file'], paths['selection'].parent/selection['quarantine_record'],
                paths['curation'], *(paths['curation'].parent/n for n in ['measurement_decisions.csv', 'curated_structures.csv', 'structure_summary.csv', 'review.json']),
                paths['preparation']]
    table('Evidence Files', 'The source files behind every table, relative to this workbook. SHA-256 identifies their exact bytes.',
          ['File', 'Bytes', 'SHA-256'],
          [[Path(os.path.relpath(p, parent)).as_posix(), p.stat().st_size, digest(p)] for p in evidence], [80, 14, 72],
          links=lambda row, i: row[0] if i == 0 else None)
    # Sanity checks tie the workbook back to the saved stage counts.
    assert len(kept) == curation['counts']['retained_activities']
    assert len(kept) + len(removed) == acquisition['activity_records'] == len(activities)
    assert len(documents) == acquisition['document_records'] and len(assays) == acquisition['assay_records']
    book.properties.creator = 'scripts/build_data_summary.py'
    raw = io.BytesIO(); book.save(raw)
    return normalize_xlsx(raw.getvalue(), curation['completed_at_utc'][:19] + 'Z')


def is_current(root=ROOT):
    """True when the saved workbook exactly matches a fresh build from current data."""
    path = Path(root)/WORKBOOK
    return path.is_file() and path.read_bytes() == render(root)


def build_summary(root=ROOT):
    root = Path(root); destination = root/WORKBOOK; temporary = destination.with_name('.~' + destination.name)
    data = render(root)
    temporary.write_bytes(data)
    # Reopen the finished file before replacing the reader's copy.
    check = load_workbook(temporary, read_only=True)
    assert check.sheetnames[0] == 'Overview' and 'Evidence Files' in check.sheetnames
    check.close(); temporary.replace(destination)
    return {'workbook': WORKBOOK.as_posix(), 'bytes': len(data), 'sha256': digest(destination)}


if __name__ == '__main__':
    print(json.dumps(build_summary(), indent=2))

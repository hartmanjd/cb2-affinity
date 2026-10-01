"""Check the assistant's audit workbook matches the saved test results it is built from."""
from pathlib import Path
import sys
import unittest

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_assistant_audit as audit  # noqa: E402


class AuditWorkbookTests(unittest.TestCase):
    def test_saved_workbook_is_current(self):
        # Rebuilding from the saved results must give the committed file byte for byte.
        self.assertEqual((ROOT / audit.WORKBOOK).read_bytes(), audit.render())

    def test_every_question_and_score_is_in_it(self):
        book = load_workbook(ROOT / audit.WORKBOOK, read_only=True)
        dataset = pd.read_csv(ROOT / "results/assistant_evaluation.csv")
        chemistry = pd.read_csv(ROOT / "results/chemistry_evaluation.csv")
        for sheet, frame in [("Dataset questions", dataset), ("Chemistry questions", chemistry)]:
            rows = list(book[sheet].iter_rows(min_row=2, values_only=True))
            self.assertEqual(len(rows), len(frame))
            outcomes = [row[6] for row in rows]
            self.assertEqual(outcomes.count("correct"), int(frame["correct"].sum()))
        overview = [row for row in book["Overview"].iter_rows(values_only=True) if row[0] in ("Dataset questions", "Chemistry questions")]
        scores = {}
        for row in overview:   # The score table comes first; the sheet guide below reuses the same names.
            scores.setdefault(row[0], row[1])
        self.assertEqual(scores,
                         {"Dataset questions": f"{int(dataset['correct'].sum())}/{len(dataset)}",
                          "Chemistry questions": f"{int(chemistry['correct'].sum())}/{len(chemistry)}"})


if __name__ == "__main__":
    unittest.main()

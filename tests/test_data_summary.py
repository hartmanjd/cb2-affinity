"""Check that the data provenance workbook is written with stable bytes."""
from pathlib import Path
import io
import sys
import time
import unittest
sys.path.insert(0, str(Path(__file__).resolve().parents[1]/'scripts'))
from openpyxl import Workbook, load_workbook
from build_data_summary import normalize_xlsx


def saved_workbook(value):
    book = Workbook()
    book.active['A1'] = value
    stream = io.BytesIO(); book.save(stream)
    return stream.getvalue()


class StableWorkbookTests(unittest.TestCase):
    def test_same_content_saved_at_different_times_gives_identical_bytes(self):
        first = saved_workbook('pKi')
        time.sleep(2)  # Zip entries record time to two-second precision.
        second = saved_workbook('pKi')
        stamp = '2026-01-01T00:00:00Z'
        self.assertEqual(normalize_xlsx(first, stamp), normalize_xlsx(second, stamp))

    def test_normalizing_keeps_the_content_and_sets_the_stamp(self):
        data = normalize_xlsx(saved_workbook('CB2'), '2026-01-01T00:00:00Z')
        book = load_workbook(io.BytesIO(data))
        self.assertEqual(book.active['A1'].value, 'CB2')
        self.assertEqual(book.properties.modified.year, 2026)

    def test_different_content_still_differs(self):
        stamp = '2026-01-01T00:00:00Z'
        self.assertNotEqual(normalize_xlsx(saved_workbook('a'), stamp), normalize_xlsx(saved_workbook('b'), stamp))
